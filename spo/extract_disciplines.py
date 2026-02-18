# -*- coding: utf-8 -*-
"""Извлекает все уникальные названия дисциплин из spo/output/_all.xlsx и сохраняет в Excel."""
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook, Workbook
except ImportError:
    print("Установите openpyxl: pip install openpyxl")
    sys.exit(1)


def clean_discipline_name(name: str) -> str:
    """Убирает из названия дисциплины суффиксы типа '1 П', '1П', '1п', '2 П', '2П', '2п' и т.д."""
    if not name:
        return ''
    # Убираем паттерны типа "1 П", "1П", "1п", "2 П", "2П", "2п", "3Б", "1Б" и т.д. в конце строки
    cleaned = re.sub(r'\s*\d+\s*[ПпБб]\s*$', '', name.strip())
    return cleaned.strip()

OUTPUT_DIR = Path(__file__).resolve().parent / "output"
EXCEL_FILE = OUTPUT_DIR / "_all.xlsx"
OUTPUT_EXCEL = OUTPUT_DIR / "disciplines_unique.xlsx"

if not EXCEL_FILE.exists():
    print(f"Файл не найден: {EXCEL_FILE}", file=sys.stderr)
    sys.exit(1)

print(f"Читаю {EXCEL_FILE}...", file=sys.stderr, flush=True)
wb = load_workbook(EXCEL_FILE, data_only=True)
ws = wb.active

# Находим колонку с дисциплинами (subject_name)
headers = []
header_row = None
for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
    if row and any(v and isinstance(v, str) and 'subject_name' in str(v).lower() for v in row):
        headers = [str(v).strip() if v else '' for v in row]
        header_row = row_idx
        break

if not headers:
    print("Не найдена строка заголовков с 'subject_name'", file=sys.stderr)
    sys.exit(1)

# Находим индекс колонки subject_name
subject_col_idx = None
for idx, h in enumerate(headers, 1):
    if h.lower() == 'subject_name':
        subject_col_idx = idx
        break

if not subject_col_idx:
    print("Не найдена колонка 'subject_name'", file=sys.stderr)
    sys.exit(1)

# Собираем все уникальные дисциплины (с очисткой от суффиксов)
disciplines = set()
data_start_row = header_row + 1

for row in ws.iter_rows(min_row=data_start_row, values_only=True):
    if not row:
        continue
    if subject_col_idx <= len(row):
        discipline = row[subject_col_idx - 1]
        if discipline and isinstance(discipline, str):
            d_clean = clean_discipline_name(discipline.strip())
            if d_clean:
                disciplines.add(d_clean)

# Сортируем и сохраняем в Excel
disciplines_sorted = sorted(disciplines)

wb = Workbook()
ws = wb.active
if ws is not None:
    ws.title = "Дисциплины СПО"
    ws.append(['Название дисциплины'])
    for disc in disciplines_sorted:
        ws.append([disc])

wb.save(OUTPUT_EXCEL)
print(f"Найдено уникальных дисциплин: {len(disciplines_sorted)}", file=sys.stderr)
print(f"Сохранено в: {OUTPUT_EXCEL}", file=sys.stderr)
